"""Read proposal files. Extraction of meaning lives in extract.py."""

from __future__ import annotations

from pathlib import Path

from brochure_agent.models import SourceDocument


class ReadError(ValueError):
    pass


def read_source(path: Path, image_dir: Path, *, ocr: bool = False) -> SourceDocument:
    if not path.is_file():
        raise ReadError(f"Proposal file not found: {path}")
    suffix = path.suffix.lower()
    image_dir.mkdir(parents=True, exist_ok=True)
    if suffix == ".txt":
        text = path.read_text(encoding="utf-8", errors="replace")
        return SourceDocument(path=str(path), text=text)
    if suffix == ".docx":
        return _read_docx(path, image_dir)
    if suffix == ".pdf":
        return _read_pdf(path, image_dir, ocr=ocr)
    raise ReadError("Proposal files must be .txt, .docx, or .pdf.")


def _read_docx(path: Path, image_dir: Path) -> SourceDocument:
    try:
        from docx import Document
    except ImportError as exc:
        raise ReadError("python-docx is not installed.") from exc
    document = Document(str(path))
    paragraphs = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
    images = _docx_images(document, image_dir)
    warnings = []
    if not paragraphs:
        warnings.append("The DOCX file has no paragraph text.")
    return SourceDocument(
        path=str(path),
        text="\n".join(paragraphs),
        images=images,
        warnings=warnings,
    )


def _docx_images(document, image_dir: Path) -> list[str]:
    saved: list[str] = []
    index = 0
    for rel in document.part.rels.values():
        if "image" not in rel.reltype:
            continue
        index += 1
        part = rel.target_part
        suffix = Path(str(part.partname)).suffix or ".bin"
        destination = image_dir / f"docx-image-{index}{suffix}"
        destination.write_bytes(part.blob)
        saved.append(str(destination))
    return saved


def _read_pdf(path: Path, image_dir: Path, *, ocr: bool) -> SourceDocument:
    try:
        import fitz
    except ImportError as exc:
        raise ReadError("PyMuPDF is not installed.") from exc
    warnings: list[str] = []
    images: list[str] = []
    pages: list[str] = []
    ocr_used = False
    with fitz.open(path) as document:
        for page_index, page in enumerate(document):
            pages.append(page.get_text() or "")
            images.extend(_pdf_images(document, page, page_index, image_dir))
        text = "\n".join(pages).strip()
        if len(text) < 30 and ocr:
            ocr_text, ocr_used, ocr_warning = _ocr_pdf(document)
            if ocr_warning:
                warnings.append(ocr_warning)
            if ocr_text:
                text = ocr_text
        elif len(text) < 30:
            warnings.append(
                "The PDF has almost no extractable text. Re-run with --ocr if it is scanned."
            )
    return SourceDocument(
        path=str(path),
        text=text,
        images=images,
        ocr_used=ocr_used,
        warnings=warnings,
    )


def _pdf_images(document, page, page_index: int, image_dir: Path) -> list[str]:
    import fitz

    saved: list[str] = []
    for image in page.get_images(full=True):
        xref = image[0]
        try:
            pixmap = fitz.Pixmap(document, xref)
            if pixmap.n - pixmap.alpha > 3:
                pixmap = fitz.Pixmap(fitz.csRGB, pixmap)
            destination = image_dir / f"pdf-page{page_index + 1}-xref{xref}.png"
            pixmap.save(str(destination))
            saved.append(str(destination))
        except Exception:
            continue
    return saved


def _ocr_pdf(document) -> tuple[str, bool, str | None]:
    try:
        import pytesseract
        from PIL import Image
    except ImportError:
        return "", False, "OCR was requested, but pytesseract or Pillow is not installed."
    import fitz

    chunks: list[str] = []
    try:
        for page in document:
            pixmap = page.get_pixmap(matrix=fitz.Matrix(2, 2), alpha=False)
            image = Image.frombytes("RGB", (pixmap.width, pixmap.height), pixmap.samples)
            chunks.append(pytesseract.image_to_string(image))
    except Exception as exc:
        return "", False, f"OCR failed: {exc}"
    text = "\n".join(chunks).strip()
    if not text:
        return "", True, "OCR ran and returned no text."
    return text, True, None
